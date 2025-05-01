import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import json
import streamlit as st
import os

# Load configuration
try:
    with open("config.json", "r") as f:
        config = json.load(f)
    PRIMARY_COLOR = config.get("primary_color", "blue")
    HEADER_FILL = PatternFill(start_color=config.get("header_fill_color", "DDDDDD"),
                               end_color=config.get("header_fill_color", "DDDDDD"),
                               fill_type="solid")
    HEADER_FONT = Font(bold=True, color=config.get("header_font_color", "000000"))
    ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
except FileNotFoundError:
    PRIMARY_COLOR = "blue"
    HEADER_FILL = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="000000")
    ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
    st.error("config.json not found. Using default styling.")

def format_header(cell):
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = ALIGNMENT

def create_financial_summary(census_df, rates_df, benefit_summary):
    """
    Creates the financial summary DataFrame.
    """
    summary_df = pd.DataFrame()
    summary_df['Plan'] = benefit_summary['Plan']
    summary_df['Employee Count'] = census_df.groupby('Plan')['Employee'].count().reindex(benefit_summary['Plan'], fill_value=0)
    summary_df = pd.merge(summary_df, rates_df, on='Plan', how='left')
    summary_df['Total Premium'] = summary_df['Employee Count'] * summary_df['Rate']
    return summary_df

def generate_excel_workbook(census_file, rates_file, benefit_summary_data):
    """
    Generates an Excel workbook with census data, rates, and financial summary.
    """
    try:
        census_df = pd.read_excel(census_file)
        rates_df = pd.read_excel(rates_file)
        benefit_summary = pd.DataFrame(benefit_summary_data)

        workbook = Workbook()

        # Sheet 1: Census Data
        census_sheet = workbook.active
        census_sheet.title = "Census Data"
        for r_idx, row in enumerate(dataframe_to_rows(census_df, header=True, index=False)):
            census_sheet.append(row)
            if r_idx == 0:
                for cell in census_sheet[1]:
                    format_header(cell)

        # Sheet 2: Rates Data
        rates_sheet = workbook.create_sheet(title="Rates Data")
        for r_idx, row in enumerate(dataframe_to_rows(rates_df, header=True, index=False)):
            rates_sheet.append(row)
            if r_idx == 0:
                for cell in rates_sheet[1]:
                    format_header(cell)

        # Sheet 3: Benefit Summary
        summary_sheet = workbook.create_sheet(title="Benefit Summary")
        for r_idx, row in enumerate(dataframe_to_rows(benefit_summary, header=True, index=False)):
            summary_sheet.append(row)
            if r_idx == 0:
                for cell in summary_sheet[1]:
                    format_header(cell)

        # Sheet 4: Financial Summary
        financial_summary_df = create_financial_summary(census_df, rates_df, benefit_summary)
        financial_sheet = workbook.create_sheet(title="Financial Summary")
        for r_idx, row in enumerate(dataframe_to_rows(financial_summary_df, header=True, index=False)):
            financial_sheet.append(row)
            if r_idx == 0:
                for cell in financial_sheet[1]:
                    format_header(cell)

        return workbook
    except FileNotFoundError as e:
        st.error(f"Error: One of the uploaded files was not found. {e}")
        return None
    except Exception as e:
        st.error(f"An error occurred during workbook generation: {e}")
        return None

if __name__ == "__main__":
    st.title("Excel Workbook Generator")

    census_file = st.file_uploader("Upload Census Data (Excel)", type=["xlsx"])
    rates_file = st.file_uploader("Upload Rates Data (Excel)", type=["xlsx"])
    benefit_summary_text = st.text_area("Enter Benefit Summary (JSON)", '[\n  {"Plan": "Plan A"}, \n  {"Plan": "Plan B"}\n]')

    if census_file and rates_file and benefit_summary_text:
        try:
            benefit_summary_data = json.loads(benefit_summary_text)
            if not isinstance(benefit_summary_data, list):
                st.error("Benefit Summary must be a JSON list of dictionaries.")
            else:
                workbook = generate_excel_workbook(census_file, rates_file, benefit_summary_data)
                if workbook:
                    st.success("Excel workbook generated successfully!")
                    # Placeholder for download button in a Streamlit context
                    # To be used within the Streamlit app (newapp.py)
                    pass
        except json.JSONDecodeError:
            st.error("Invalid JSON format for Benefit Summary.")